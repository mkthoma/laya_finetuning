"""Plain tables and their .xlsx writer (openpyxl) for the Phase 7 results workbook.

A Sheet stacks Tables vertically. The first table's header is row 1 (bold, frozen, auto-filtered); each later table
follows one blank row with a bold title row and its own header. Cells keep their type: numbers stay numbers with the
column's number format, booleans stay booleans, None is an empty cell, a list of scalars becomes "a, b". The writer
refuses any cell that looks like a data row id (`<split>-NNNNNN`): the workbook is metrics only (design Appendix D),
and flatten_records (for inputs of unknown layout) drops record-level keys and row-id-like strings before that.
"""
from __future__ import annotations

import json
import math
import os
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence

ROW_ID = re.compile(r"\b[a-z][a-z0-9_]*-\d{6}\b")
# An absolute local path (a Windows drive, or a POSIX home / Colab / temp root): may name the user; URLs never match.
LOCAL_PATH = re.compile(r"(?<![A-Za-z0-9])[A-Za-z]:[\\/]"
                        r"|(?:^|[\s(\"'=`])/(?:home|Users|content|root|mnt|tmp|var|opt)/")
F4, F3, F2, F1, F0, PCT, SCI = "0.0000", "0.000", "0.00", "0.0", "0", "0.0%", "0.00E+00"
MIN_WIDTH, MAX_WIDTH, WRAP_WIDTH = 8, 60, 110
LINE_HEIGHT, MAX_ROW_HEIGHT = 15, 409          # points; 409 is the Excel maximum
HEADER_FILL = "DDEBF7"
TOP_LEVEL = "(top level)"
WITHHELD_KEYS = frozenset({"id", "ids", "fsq_place_id", "fsq_place_ids", "name", "names", "state", "states", "text",
                           "row", "rows", "record", "records", "address", "preds", "path", "paths"})


@dataclass(frozen=True)
class Column:
    key: str
    header: str
    fmt: str | None = None


@dataclass(frozen=True)
class Table:
    title: str | None
    columns: Sequence[Column]
    rows: Sequence[Mapping[str, Any]]


@dataclass(frozen=True)
class Sheet:
    name: str
    tables: Sequence[Table]
    wrap: bool = False


def cols(*specs: tuple) -> list[Column]:
    """Column specs as tuples: (key, header) or (key, header, fmt)."""
    return [Column(*s) for s in specs]


def kv_table(title: str | None, pairs: Iterable[tuple], headers: Sequence[str] = ("Item", "Value")) -> Table:
    """A two- (or three-) column item / value (/ source) table."""
    keys = ("item", "value", "source")[:len(headers)]
    return Table(title, [Column(k, h) for k, h in zip(keys, headers)], [dict(zip(keys, p)) for p in pairs])


def message_table(text: str, header: str = "Status") -> Table:
    return Table(None, [Column("status", header)], [{"status": text}])


def cell_value(v: Any) -> Any:
    """A value openpyxl can store, keeping numbers and booleans as they are."""
    if isinstance(v, float) and not math.isfinite(v):
        return None
    if v is None or isinstance(v, (bool, int, float, str)):
        return v
    if isinstance(v, (list, tuple)) and all(not isinstance(x, (Mapping, list, tuple)) for x in v):
        return ", ".join(str(x) for x in v)
    return json.dumps(v, ensure_ascii=False, sort_keys=True, default=str)


# ---------------------------------------------------------------- inputs of unknown layout

def _nested(v: Any) -> bool:
    return isinstance(v, Mapping) or (isinstance(v, list) and any(isinstance(x, (Mapping, list)) for x in v))


def flatten_records(obj: Any) -> tuple[list[dict[str, Any]], int]:
    """Every dict with scalar leaves -> {"Path": "a / b", **leaves} (nested dicts and lists of dicts recurse; lists
    of scalars are joined). Keys naming records (WITHHELD_KEYS) and row-id-like keys or strings are dropped; the
    second value counts what was dropped."""
    out: list[dict[str, Any]] = []
    dropped = 0

    def walk(node: Any, path: tuple[str, ...]) -> None:
        nonlocal dropped
        items = node.items() if isinstance(node, Mapping) else ((f"[{i}]", x) for i, x in enumerate(node))
        leaves, deeper = {}, []
        for k, v in items:
            k = str(k)
            if k.lower() in WITHHELD_KEYS or ROW_ID.search(k):
                dropped += 1
            elif _nested(v):
                deeper.append((k, v))
            elif isinstance(cv := cell_value(v), str) and (ROW_ID.search(cv) or LOCAL_PATH.search(cv)):
                dropped += 1
            else:
                leaves[k] = cv
        if leaves:
            out.append({"Path": " / ".join(path) or TOP_LEVEL, **leaves})
        for k, v in deeper:
            walk(v, (*path, k))

    if _nested(obj):
        walk(obj, ())
    return out, dropped


def records_table(title: str | None, records: Sequence[Mapping[str, Any]]) -> Table:
    """A table over the union of the records' keys (first-seen order)."""
    keys = list(dict.fromkeys(k for r in records for k in r))
    return Table(title, [Column(k, k) for k in keys], list(records))


# ---------------------------------------------------------------- checks and writing

def check_metrics_only(sheets: Sequence[Sheet]) -> None:
    """Raise if any title, header or cell looks like a data row id or an absolute local path."""
    for s in sheets:
        for t in s.tables:
            texts = [t.title or "", *(c.header for c in t.columns),
                     *(v for r in t.rows for c in t.columns if isinstance(v := cell_value(r.get(c.key)), str))]
            hit = next((m.group(0) for x in texts if (m := ROW_ID.search(x))), None)
            if hit:
                raise ValueError(f"sheet {s.name!r}: a cell holds a row id ({hit}); the workbook is metrics only")
            if any(LOCAL_PATH.search(x) for x in texts):
                raise ValueError(f"sheet {s.name!r}: a cell holds an absolute local path; the workbook is public "
                                 f"(use repo-relative paths)")


def _shown_len(v: Any) -> int:
    if isinstance(v, float):
        return len(f"{v:.4f}")
    return max((len(line) for line in str(v).splitlines()), default=0) if v is not None else 0


def _write_table(ws: Any, t: Table, top: int, first: bool, wrap: bool, widths: dict[int, int]) -> int:
    """Write one table from row `top`; returns the next free row."""
    from openpyxl.styles import Alignment, Font, PatternFill

    r = top
    if t.title and not first:
        ws.cell(row=r, column=1, value=t.title).font = Font(bold=True, size=12)
        r += 1
    for j, c in enumerate(t.columns, 1):
        cell = ws.cell(row=r, column=j, value=c.header)
        cell.font, cell.fill = Font(bold=True), PatternFill("solid", fgColor=HEADER_FILL)
        widths[j] = max(widths.get(j, 0), len(c.header))
    head = r
    for row in t.rows:
        r += 1
        for j, c in enumerate(t.columns, 1):
            v = cell_value(row.get(c.key))
            cell = ws.cell(row=r, column=j, value=v)
            if c.fmt and isinstance(v, (int, float)) and not isinstance(v, bool):
                cell.number_format = c.fmt
            if wrap:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            widths[j] = max(widths.get(j, 0), _shown_len(v))
    if first:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A{head}:{ws.cell(row=head, column=len(t.columns)).column_letter}{max(r, head)}"
    return r + 1


def _fit_heights(ws: Any) -> None:
    """Row heights for wrapped text (Excel does not reliably auto-fit rows of a generated file)."""
    for row in ws.iter_rows():
        lines = 1
        for cell in row:
            if isinstance(cell.value, str):
                width = max(1, int(ws.column_dimensions[cell.column_letter].width or MIN_WIDTH) - 2)
                lines = max(lines, sum(max(1, math.ceil(len(ln) / width)) for ln in cell.value.splitlines()))
        if lines > 1:
            ws.row_dimensions[row[0].row].height = min(LINE_HEIGHT * lines, MAX_ROW_HEIGHT)


def write_sheet(ws: Any, sheet: Sheet) -> None:
    from openpyxl.utils import get_column_letter

    widths: dict[int, int] = {}
    r = 1
    for i, t in enumerate(sheet.tables):
        r = _write_table(ws, t, r + (1 if i else 0), i == 0, sheet.wrap, widths)
    for j, w in widths.items():
        cap = WRAP_WIDTH if sheet.wrap else MAX_WIDTH
        ws.column_dimensions[get_column_letter(j)].width = max(MIN_WIDTH, min(w + 2, cap))
    if sheet.wrap:
        _fit_heights(ws)


def write_workbook(sheets: Sequence[Sheet], out: Path, creator: str) -> Path:
    """Check (metrics only), then write the .xlsx atomically (temp file, then rename)."""
    from openpyxl import Workbook

    check_metrics_only(sheets)
    wb = Workbook()
    wb.remove(wb.active)
    for s in sheets:
        write_sheet(wb.create_sheet(s.name), s)
    wb.properties.creator = creator
    out.parent.mkdir(parents=True, exist_ok=True)
    tmp = out.with_name(out.name + ".tmp")
    wb.save(tmp)
    os.replace(tmp, out)
    return out
